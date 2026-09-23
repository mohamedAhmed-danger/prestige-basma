import csv
import io
from datetime import timezone
from zoneinfo import ZoneInfo

from flask import Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import User, Attendance

CAIRO_TZ = ZoneInfo("Africa/Cairo")


def _to_cairo(value, fmt='%Y-%m-%d %I:%M %p'):
    """Same display-only UTC -> Africa/Cairo rule as the cairo_time Jinja filter,
    used here for the export functions instead of a template."""
    if value is None:
        return ''
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(CAIRO_TZ).strftime(fmt)


def get_report(date_from, date_to, zone_id=None, status=None, employee_id=None):
    query = (
        Attendance.query
        .join(User, Attendance.user_id == User.id)
        .filter(Attendance.date >= date_from, Attendance.date <= date_to)
    )

    if zone_id:
        query = query.filter(Attendance.zone_id == zone_id)

    if status == 'full_shift':
        query = query.filter(Attendance.shift_status == 'full_shift')
    elif status == 'short_shift':
        query = query.filter(Attendance.shift_status == 'short_shift')
    elif status == 'active':
        query = query.filter(Attendance.status == 'active')

    if employee_id:
        query = query.filter(Attendance.user_id == employee_id)

    return query.order_by(Attendance.date.desc(), Attendance.sign_in_time.desc()).all()


STATUS_LABELS = {
    'full_shift': 'شفت كامل',
    'short_shift': 'شفت غير كامل',
    None: 'قيد العمل (لم ينصرف)',
}


def export_csv(rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([
        'التاريخ', 'الموظف', 'الفرع', 'وقت الحضور', 'وقت الانصراف',
        'إجمالي الساعات', 'حالة الشفت', 'متأخر', 'دقائق التأخير',
    ])

    for row in rows:
        writer.writerow([
            row.date.isoformat(),
            row.user.name,
            row.zone.name,
            _to_cairo(row.sign_in_time),
            _to_cairo(row.sign_out_time),
            f'{row.total_hours:.1f}' if row.total_hours is not None else '',
            STATUS_LABELS.get(row.shift_status, row.shift_status),
            'نعم' if row.is_late else 'لا',
            row.late_minutes or 0,
        ])

    # UTF-8 BOM so Excel opens the Arabic columns correctly instead of mojibake.
    csv_bytes = '\ufeff' + buffer.getvalue()

    return Response(
        csv_bytes,
        mimetype='text/csv',
        headers={'Content-Disposition': 'attachment; filename=attendance_report.csv'},
    )


def export_excel(rows):
    """Generates a professionally formatted Excel (.xlsx) file with RTL layout,
    custom headers, status badge highlighting, and column width auto-fitting."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "تقرير الحضور والغياب"
    ws.views.sheetView[0].rightToLeft = True
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10, color="0F172A")
    
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_row_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    fill_ok = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_ok = Font(name="Segoe UI", size=10, bold=True, color="166534")

    fill_warn = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    font_warn = Font(name="Segoe UI", size=10, bold=True, color="854D0E")

    fill_info = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    font_info = Font(name="Segoe UI", size=10, bold=True, color="1E40AF")

    fill_danger = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_danger = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Sheet Title Banner
    ws.merge_cells('A1:I1')
    ws['A1'] = "تقرير حضور وانصراف الموظفين — برستيج بصمة"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_right
    ws.row_dimensions[1].height = 35

    headers = [
        'التاريخ', 'اسم الموظف', 'الفرع', 'وقت الحضور', 'وقت الانصراف',
        'إجمالي الساعات', 'حالة الشفت', 'التأخير', 'دقائق التأخير'
    ]

    ws.append([])  # Blank row 2

    # Append Header Row at row 3
    ws.append(headers)
    ws.row_dimensions[3].height = 26
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # Append Data Rows
    current_row = 4
    for row in rows:
        formatted_hours = f'{row.total_hours:.1f} ساعة' if row.total_hours is not None else '—'
        shift_status_text = STATUS_LABELS.get(row.shift_status, row.shift_status)
        late_text = 'نعم' if row.is_late else 'لا'

        data = [
            row.date.isoformat(),
            row.user.name,
            row.zone.name,
            _to_cairo(row.sign_in_time),
            _to_cairo(row.sign_out_time) if row.sign_out_time else '—',
            formatted_hours,
            shift_status_text,
            late_text,
            row.late_minutes or 0,
        ]
        ws.append(data)
        ws.row_dimensions[current_row].height = 22

        for col_num in range(1, len(data) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = font_body
            cell.alignment = align_center
            cell.border = thin_border
            if current_row % 2 == 1:
                cell.fill = fill_row_alt

            # Status Column (Column 7) styling
            if col_num == 7:
                if row.shift_status == 'full_shift':
                    cell.fill = fill_ok
                    cell.font = font_ok
                elif row.shift_status == 'short_shift':
                    cell.fill = fill_warn
                    cell.font = font_warn
                elif row.status == 'active':
                    cell.fill = fill_info
                    cell.font = font_info

            # Late Column (Column 8) styling
            if col_num == 8 and row.is_late:
                cell.fill = fill_danger
                cell.font = font_danger

        current_row += 1

    # Auto-adjust column widths cleanly
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 6, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=attendance_report.xlsx'},
    )
